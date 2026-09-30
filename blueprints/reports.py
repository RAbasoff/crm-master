"""
reports blueprint
"""
from datetime import datetime, timedelta
from flask import (Blueprint, render_template, request, redirect, url_for, flash,
                   jsonify, send_file, session, g)
from flask_login import login_required, current_user
from flask_babel import gettext as _
from werkzeug.utils import secure_filename
import os, io, json

from models import (db, AuditLog, CylinderLog, CylinderOrder, FactorySection, FaultReport, GasCylinder, Machine, Monteur, Opdracht, SystemLog, User, UserActivityLog, Verantwoordelijke, VoorraadMutatie)
from utils import role_required, safe_date, safe_int

bp = Blueprint('reports', __name__)

@bp.route('/reports')
@login_required
@role_required('admin', 'director', 'technician')
def reports():
    return render_template('reports.html')


@bp.route('/reports/advanced')
@login_required
@role_required('admin', 'director')
def reports_advanced():
    report_type = request.args.get('type', 'activity')
    date_from = request.args.get('date_from', (datetime.utcnow() - timedelta(days=30)).strftime('%Y-%m-%d'))
    date_to = request.args.get('date_to', datetime.utcnow().strftime('%Y-%m-%d'))
    user_id = request.args.get('user_id', '')
    section_id = request.args.get('section_id', '')
    
    d_from = safe_date(date_from) or (datetime.utcnow() - timedelta(days=30))
    d_to = (safe_date(date_to) or datetime.utcnow()) + timedelta(days=1)
    
    users = User.query.filter(User.is_active_user == True).order_by(User.display_name).all()
    sections = FactorySection.query.order_by(FactorySection.name).all()
    
    data = []
    stats = {}
    
    if report_type == 'activity':
        q = UserActivityLog.query.filter(UserActivityLog.created_at >= d_from, UserActivityLog.created_at < d_to)
        if user_id:
            q = q.filter_by(user_id=safe_int(user_id))
        data = q.order_by(UserActivityLog.created_at.desc()).limit(500).all()
        stats['total_actions'] = q.count()
        stats['unique_users'] = db.session.query(db.func.count(db.distinct(UserActivityLog.user_id))).filter(UserActivityLog.created_at >= d_from, UserActivityLog.created_at < d_to).scalar()
        
    elif report_type == 'faults':
        q = FaultReport.query.filter(FaultReport.created_at >= d_from, FaultReport.created_at < d_to)
        if user_id:
            q = q.filter_by(reporter_id=safe_int(user_id))
        if section_id:
            q = q.filter(FaultReport.machine.has(Machine.section_id == safe_int(section_id)))
        data = q.order_by(FaultReport.created_at.desc()).all()
        stats['total'] = q.count()
        stats['open'] = q.filter(FaultReport.status.in_(['open', 'accepted', 'in_progress'])).count()
        stats['resolved'] = q.filter_by(status='resolved').count()
        stats['critical'] = q.filter_by(priority='critical').count()
        
    elif report_type == 'warehouse':
        q = VoorraadMutatie.query.filter(VoorraadMutatie.aangemaakt >= d_from, VoorraadMutatie.aangemaakt < d_to)
        data = q.order_by(VoorraadMutatie.aangemaakt.desc()).limit(500).all()
        stats['total_movements'] = q.count()
        stats['incoming'] = q.filter_by(type='inkomend').count()
        stats['outgoing'] = q.filter_by(type='uitgaand').count()
        
    elif report_type == 'errors':
        q = SystemLog.query.filter(SystemLog.created_at >= d_from, SystemLog.created_at < d_to)
        if user_id:
            q = q.filter_by(user_id=safe_int(user_id))
        data = q.order_by(SystemLog.created_at.desc()).limit(500).all()
        stats['total'] = q.count()
        stats['errors'] = q.filter_by(level='ERROR').count()
        stats['warnings'] = q.filter_by(level='WARNING').count()
        
    elif report_type == 'users':
        q = AuditLog.query.filter(AuditLog.created_at >= d_from, AuditLog.created_at < d_to)
        if user_id:
            q = q.filter_by(user_id=safe_int(user_id))
        data = q.order_by(AuditLog.created_at.desc()).limit(500).all()
        stats['total'] = q.count()
        
    elif report_type == 'responsible':
        # Report on responsible persons - who submitted what and when
        persons = Verantwoordelijke.query.filter(Verantwoordelijke.is_active == True).order_by(Verantwoordelijke.naam).all()
        responsible_data = []
        for p in persons:
            faults = FaultReport.query.filter(
                FaultReport.reporter_id.in_(
                    db.session.query(User.id).filter(User.person_id == p.id)
                ),
                FaultReport.created_at >= d_from,
                FaultReport.created_at < d_to
            ).order_by(FaultReport.created_at.desc()).all()
            if faults or not user_id:
                responsible_data.append({
                    'person': p,
                    'faults': faults,
                    'total': len(faults),
                    'open': len([f for f in faults if f.status in ['open', 'accepted', 'in_progress']]),
                    'resolved': len([f for f in faults if f.status == 'resolved']),
                })
        data = responsible_data
        stats['total_persons'] = len(responsible_data)
        stats['total_faults'] = sum(r['total'] for r in responsible_data)
        
    elif report_type == 'machines':
        # Report by machines - faults, maintenance, status
        machines_q = Machine.query.order_by(Machine.name)
        if section_id:
            machines_q = machines_q.filter_by(section_id=safe_int(section_id))
        machines_list = machines_q.all()
        machine_data = []
        for m in machines_list:
            faults = FaultReport.query.filter(
                FaultReport.machine_id == m.id,
                FaultReport.created_at >= d_from,
                FaultReport.created_at < d_to
            ).order_by(FaultReport.created_at.desc()).all()
            total_faults = FaultReport.query.filter(FaultReport.machine_id == m.id).count()
            machine_data.append({
                'machine': m,
                'faults': faults,
                'period_count': len(faults),
                'total_count': total_faults,
                'open': len([f for f in faults if f.status in ['open', 'accepted', 'in_progress']]),
                'critical': len([f for f in faults if f.priority == 'critical']),
            })
        data = machine_data
        stats['total_machines'] = len(machine_data)
        stats['total_faults'] = sum(r['period_count'] for r in machine_data)
        stats['machines_with_faults'] = len([r for r in machine_data if r['period_count'] > 0])

    elif report_type == 'gas':
        from models import GasCylinder, CylinderLog, CylinderOrder
        # Received: cylinders created in period
        received_logs = CylinderLog.query.filter(
            CylinderLog.action == 'created',
            CylinderLog.date >= d_from, CylinderLog.date < d_to
        ).all()
        received_ids = list(set(l.cylinder_id for l in received_logs if l.cylinder_id))
        n2_received = sum(1 for cid in received_ids if GasCylinder.query.get(cid) and GasCylinder.query.get(cid).gas_type == 'nitrogen')
        co2_received = sum(1 for cid in received_ids if GasCylinder.query.get(cid) and GasCylinder.query.get(cid).gas_type == 'co2')

        # Consumed: cylinders that became empty in period
        consumed_logs = CylinderLog.query.filter(
            CylinderLog.action.like('%_to_empty%'),
            CylinderLog.date >= d_from, CylinderLog.date < d_to
        ).all()
        consumed_ids = list(set(l.cylinder_id for l in consumed_logs if l.cylinder_id))
        n2_consumed = sum(1 for cid in consumed_ids if GasCylinder.query.get(cid) and GasCylinder.query.get(cid).gas_type == 'nitrogen')
        co2_consumed = sum(1 for cid in consumed_ids if GasCylinder.query.get(cid) and GasCylinder.query.get(cid).gas_type == 'co2')

        # Monthly breakdown (last 6 months)
        chart_data = []
        for i in range(5, -1, -1):
            m_start = (datetime.utcnow().replace(day=1) - timedelta(days=i*30)).replace(day=1)
            m_end = (m_start + timedelta(days=32)).replace(day=1)
            m_key = m_start.strftime('%Y-%m')
            r_logs = CylinderLog.query.filter(CylinderLog.action == 'created', CylinderLog.date >= m_start, CylinderLog.date < m_end).all()
            c_logs = CylinderLog.query.filter(CylinderLog.action.like('%_to_empty%'), CylinderLog.date >= m_start, CylinderLog.date < m_end).all()
            r_ids = set(l.cylinder_id for l in r_logs if l.cylinder_id)
            c_ids = set(l.cylinder_id for l in c_logs if l.cylinder_id)
            chart_data.append({
                'month': m_key,
                'received': len(r_ids),
                'consumed': len(c_ids)
            })

        # Orders in period
        orders = CylinderOrder.query.filter(
            CylinderOrder.ordered_at >= d_from, CylinderOrder.ordered_at < d_to
        ).order_by(CylinderOrder.ordered_at.desc()).all()

        # All logs in period
        all_logs = CylinderLog.query.filter(
            CylinderLog.date >= d_from, CylinderLog.date < d_to
        ).order_by(CylinderLog.date.desc()).limit(200).all()

        stats['n2_received'] = n2_received
        stats['co2_received'] = co2_received
        stats['n2_consumed'] = n2_consumed
        stats['co2_consumed'] = co2_consumed
        stats['total_received'] = n2_received + co2_received
        stats['total_consumed'] = n2_consumed + co2_consumed
        stats['total_orders'] = len(orders)
        stats['chart_data'] = chart_data
        data = all_logs

    return render_template('reports_advanced.html',
        report_type=report_type, data=data, stats=stats,
        users=users, sections=sections,
        date_from=date_from, date_to=date_to,
        user_id=user_id, section_id=section_id)


@bp.route('/reports/advanced/export')
@login_required
@role_required('admin', 'director')
def reports_export():
    report_type = request.args.get('type', 'activity')
    format_type = request.args.get('format', 'csv')
    date_from = request.args.get('date_from', (datetime.utcnow() - timedelta(days=30)).strftime('%Y-%m-%d'))
    date_to = request.args.get('date_to', datetime.utcnow().strftime('%Y-%m-%d'))
    user_id = request.args.get('user_id', '')
    
    d_from = safe_date(date_from) or (datetime.utcnow() - timedelta(days=30))
    d_to = (safe_date(date_to) or datetime.utcnow()) + timedelta(days=1)
    
    # Collect data
    headers = []
    rows = []
    title = ''
    
    if report_type == 'activity':
        title = 'Активность пользователей'
        headers = ['Дата', 'Пользователь', 'Действие', 'Страница', 'Детали', 'IP']
        q = UserActivityLog.query.filter(UserActivityLog.created_at >= d_from, UserActivityLog.created_at < d_to)
        if user_id: q = q.filter_by(user_id=safe_int(user_id))
        for r in q.order_by(UserActivityLog.created_at.desc()).limit(1000).all():
            rows.append([r.created_at.strftime('%Y-%m-%d %H:%M'), r.username or '', r.action or '', r.page or '', r.details or '', r.ip_address or ''])
    elif report_type == 'faults':
        title = 'Заявки о неисправности'
        headers = ['ID', 'Дата', 'Заголовок', 'Станок', 'Приоритет', 'Статус', 'Репортер']
        q = FaultReport.query.filter(FaultReport.created_at >= d_from, FaultReport.created_at < d_to)
        if user_id: q = q.filter_by(reporter_id=safe_int(user_id))
        for f in q.order_by(FaultReport.created_at.desc()).all():
            rows.append([str(f.id), f.created_at.strftime('%Y-%m-%d %H:%M'), f.title or '', f.machine.name if f.machine else '', f.priority or '', f.status or '', f.reporter_label])
    elif report_type == 'warehouse':
        title = 'Движение склада'
        headers = ['Дата', 'Товар', 'Тип', 'Количество', 'Комментарий']
        q = VoorraadMutatie.query.filter(VoorraadMutatie.aangemaakt >= d_from, VoorraadMutatie.aangemaakt < d_to)
        for m in q.order_by(VoorraadMutatie.aangemaakt.desc()).limit(1000).all():
            rows.append([m.aangemaakt.strftime('%Y-%m-%d %H:%M'), m.item.naam if m.item else '', m.type or '', str(m.hoeveelheid), m.opmerking or ''])
    elif report_type == 'errors':
        title = 'Ошибки и предупреждения'
        headers = ['Дата', 'Уровень', 'Категория', 'Сообщение', 'Источник', 'Пользователь']
        q = SystemLog.query.filter(SystemLog.created_at >= d_from, SystemLog.created_at < d_to)
        if user_id: q = q.filter_by(user_id=safe_int(user_id))
        for r in q.order_by(SystemLog.created_at.desc()).limit(1000).all():
            rows.append([r.created_at.strftime('%Y-%m-%d %H:%M'), r.level or '', r.category or '', r.message or '', r.source or '', r.user.display_name if r.user else ''])
    elif report_type == 'users':
        title = 'Журнал аудита'
        headers = ['Дата', 'Пользователь', 'Действие', 'Тип', 'Детали', 'IP']
        q = AuditLog.query.filter(AuditLog.created_at >= d_from, AuditLog.created_at < d_to)
        if user_id: q = q.filter_by(user_id=safe_int(user_id))
        for r in q.order_by(AuditLog.created_at.desc()).limit(1000).all():
            rows.append([r.created_at.strftime('%Y-%m-%d %H:%M'), r.user.display_name if r.user else '', r.action or '', r.entity_type or '', r.details or '', r.ip_address or ''])
    elif report_type == 'responsible':
        title = 'Отчёт по ответственным'
        headers = ['Ответственный', 'Должность', 'Телефон', 'Внутр. номер', 'Email', 'Всего заявок', 'Открытых', 'Решённых']
        persons = Verantwoordelijke.query.filter(Verantwoordelijke.is_active == True).order_by(Verantwoordelijke.naam).all()
        for p in persons:
            fault_count = FaultReport.query.filter(
                FaultReport.reporter_id.in_(db.session.query(User.id).filter(User.person_id == p.id)),
                FaultReport.created_at >= d_from, FaultReport.created_at < d_to
            ).count()
            open_count = FaultReport.query.filter(
                FaultReport.reporter_id.in_(db.session.query(User.id).filter(User.person_id == p.id)),
                FaultReport.created_at >= d_from, FaultReport.created_at < d_to,
                FaultReport.status.in_(['open', 'accepted', 'in_progress'])
            ).count()
            resolved_count = FaultReport.query.filter(
                FaultReport.reporter_id.in_(db.session.query(User.id).filter(User.person_id == p.id)),
                FaultReport.created_at >= d_from, FaultReport.created_at < d_to,
                FaultReport.status == 'resolved'
            ).count()
            rows.append([p.naam or '', p.position or '', p.telefoon or '', p.internal_phone or '', p.email or '', str(fault_count), str(open_count), str(resolved_count)])
    elif report_type == 'machines':
        title = 'Отчёт по станкам'
        headers = ['Станок', 'Тип', 'Серийный номер', 'Отдел', 'Заявок за период', 'Всего заявок', 'Открытых', 'Критичных']
        machines_q = Machine.query.order_by(Machine.name)
        if section_id: machines_q = machines_q.filter_by(section_id=safe_int(section_id))
        for m in machines_q.all():
            period_count = FaultReport.query.filter(FaultReport.machine_id == m.id, FaultReport.created_at >= d_from, FaultReport.created_at < d_to).count()
            total_count = FaultReport.query.filter(FaultReport.machine_id == m.id).count()
            open_count = FaultReport.query.filter(FaultReport.machine_id == m.id, FaultReport.status.in_(['open', 'accepted', 'in_progress'])).count()
            critical_count = FaultReport.query.filter(FaultReport.machine_id == m.id, FaultReport.priority == 'critical').count()
            rows.append([m.name or '', m.machine_type or '', m.serial_number or '', m.section.name if m.section else '', str(period_count), str(total_count), str(open_count), str(critical_count)])
    
    period = f'{date_from} — {date_to}'
    filename = f'report_{report_type}_{date_from}_{date_to}'
    
    # === CSV ===
    if format_type == 'csv':
        import csv, io
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerows(rows)
        output.seek(0)
        from flask import Response
        return Response(output.getvalue(), mimetype='text/csv',
            headers={'Content-Disposition': f'attachment;filename={filename}.csv'})
    
    # === EXCEL ===
    elif format_type == 'excel':
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        wb = Workbook()
        ws = wb.active
        ws.title = title[:31]
        
        # Title
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws['A1'] = title
        ws['A1'].font = Font(bold=True, size=14)
        ws['A2'] = f'Период: {period}'
        ws['A2'].font = Font(size=10, color='666666')
        
        # Headers
        header_fill = PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF', size=11)
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin'))
        
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
            cell.border = thin_border
        
        # Data
        for row_idx, row_data in enumerate(rows, 5):
            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.border = thin_border
                cell.alignment = Alignment(wrap_text=True)
        
        # Auto-width
        for col in range(1, len(headers) + 1):
            max_len = len(headers[col-1])
            for row in range(5, len(rows) + 5):
                cell_val = str(ws.cell(row=row, column=col).value or '')
                max_len = max(max_len, min(len(cell_val), 40))
            ws.column_dimensions[ws.cell(row=4, column=col).column_letter].width = max_len + 4
        
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            download_name=f'{filename}.xlsx', as_attachment=True)
    
    # === WORD ===
    elif format_type == 'word':
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.table import WD_TABLE_ALIGNMENT
        doc = Document()
        
        # Title
        doc.add_heading(title, level=1)
        doc.add_paragraph(f'Период: {period}')
        doc.add_paragraph(f'Сгенерировано: {datetime.utcnow().strftime("%d.%m.%Y %H:%M")}')
        
        # Table
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = 'Light Grid Accent 1'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Headers
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = h
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True
                    run.font.size = Pt(9)
        
        # Data
        for row_data in rows:
            row = table.add_row()
            for i, val in enumerate(row_data):
                row.cells[i].text = str(val)
                for paragraph in row.cells[i].paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(8)
        
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return send_file(buf, mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            download_name=f'{filename}.docx', as_attachment=True)
    
    # === PDF ===
    elif format_type == 'pdf':
        # Generate HTML table for PDF
        html = f'''<!DOCTYPE html>
<html><head><meta charset="utf-8">
<style>
body {{ font-family: Arial; font-size: 10px; padding: 20px; }}
h1 {{ font-size: 16px; margin-bottom: 5px; }}
p {{ color: #666; font-size: 11px; margin-bottom: 15px; }}
table {{ width: 100%; border-collapse: collapse; font-size: 9px; }}
th {{ background: #2c3e50; color: white; padding: 6px 8px; text-align: left; font-weight: bold; }}
td {{ padding: 5px 8px; border-bottom: 1px solid #eee; }}
tr:nth-child(even) {{ background: #f9f9f9; }}
.footer {{ margin-top: 20px; font-size: 8px; color: #999; text-align: center; }}
</style></head><body>
<h1>{title}</h1>
<p>Период: {period} | Сгенерировано: {datetime.utcnow().strftime("%d.%m.%Y %H:%M")}</p>
<table><tr>'''
        for h in headers:
            html += f'<th>{h}</th>'
        html += '</tr>'
        for row in rows:
            html += '<tr>'
            for val in row:
                html += f'<td>{val}</td>'
            html += '</tr>'
        html += f'</table><div class="footer">CRM Мастерская — {title} — {period}</div></body></html>'
        
        try:
            import pdfkit
            pdf = pdfkit.from_string(html, False)
            from flask import Response
            return Response(pdf, mimetype='application/pdf',
                headers={'Content-Disposition': f'attachment;filename={filename}.pdf'})
        except Exception:
            # Fallback: return HTML that can be printed as PDF
            from flask import Response
            return Response(html, mimetype='text/html',
                headers={'Content-Disposition': f'attachment;filename={filename}.html'})
    
    return jsonify({'error': 'Unknown format'}), 400


@bp.route('/reports/period', methods=['GET', 'POST'])
@login_required
@role_required('admin', 'director', 'technician')
def report_period():
    if request.method == 'POST':
        try:
            d_from = datetime.strptime(request.form['date_from'], '%Y-%m-%d')
            d_to = datetime.strptime(request.form['date_to'], '%Y-%m-%d') + timedelta(days=1)
        except (ValueError, KeyError):
            flash(_('Invalid date format'), 'error')
            return redirect(url_for('reports.report_period'))
    else:
        d_to = datetime.utcnow()
        d_from = d_to - timedelta(days=30)
    orders = Opdracht.query.filter(Opdracht.aangemaakt >= d_from, Opdracht.aangemaakt < d_to).order_by(Opdracht.aangemaakt.desc()).all()
    omzet = sum(o.totaal for o in orders if o.status == 'afgeleverd')
    voltooid = len([o for o in orders if o.status == 'afgeleverd'])
    geannuleerd = len([o for o in orders if o.status == 'geannuleerd'])
    # Pre-fetch all monteurs in one query
    monteur_ids = list(set(o.monteur_id for o in orders if o.monteur_id))
    monteurs_map = {m.id: m.naam for m in Monteur.query.filter(Monteur.id.in_(monteur_ids)).all()} if monteur_ids else {}
    ws = {}
    for o in orders:
        if o.monteur_id:
            if o.monteur_id not in ws:
                ws[o.monteur_id] = {'naam': monteurs_map.get(o.monteur_id, 'Onbekend'), 'orders': 0, 'omzet': 0, 'voltooid': 0}
            ws[o.monteur_id]['orders'] += 1
            if o.status == 'afgeleverd':
                ws[o.monteur_id]['omzet'] += o.totaal
                ws[o.monteur_id]['voltooid'] += 1
    return render_template('report_period.html', orders=orders,
                         date_from=d_from.strftime('%Y-%m-%d'),
                         date_to=(d_to - timedelta(days=1)).strftime('%Y-%m-%d'),
                         total_revenue=omzet, total_orders=len(orders),
                         completed=voltooid, cancelled=geannuleerd, worker_stats=ws)


@bp.route('/reports/worker/<int:worker_id>')
@login_required
@role_required('admin', 'director', 'technician')
def report_worker(worker_id):
    w = Monteur.query.get_or_404(worker_id)
    df = request.args.get('date_from', (datetime.utcnow() - timedelta(days=30)).strftime('%Y-%m-%d'))
    dt = request.args.get('date_to', datetime.utcnow().strftime('%Y-%m-%d'))
    orders = Opdracht.query.filter(Opdracht.monteur_id == worker_id,
        Opdracht.aangemaakt >= datetime.strptime(df, '%Y-%m-%d'),
        Opdracht.aangemaakt < datetime.strptime(dt, '%Y-%m-%d') + timedelta(days=1)
    ).order_by(Opdracht.aangemaakt.desc()).all()
    omzet = sum(o.totaal for o in orders if o.status == 'afgeleverd')
    voltooid = len([o for o in orders if o.status == 'afgeleverd'])
    avg = 0
    gereed = [o for o in orders if o.gereed and o.gestart]
    if gereed: avg = sum((o.gereed - o.gestart).total_seconds()/3600 for o in gereed) / len(gereed)
    return render_template('report_worker.html', worker=w, orders=orders, date_from=df, date_to=dt,
                         total_revenue=omzet, completed=voltooid, avg_time=round(avg,1))
