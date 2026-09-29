"""
assets blueprint
"""
from datetime import datetime, timedelta
from flask import (Blueprint, render_template, request, redirect, url_for, flash,
                   jsonify, send_file, session, g)
from flask_login import login_required, current_user
from flask_babel import gettext as _
from werkzeug.utils import secure_filename
import os, io, json

from models import (db, Contractor, Equipment, EquipmentServiceLog, FactorySection, User, Verantwoordelijke)
from utils import log_audit, role_required, safe_commit, safe_date, safe_float, safe_int, save_uploaded_file, ensure_fpdf, find_pdf_font

bp = Blueprint('assets', __name__)

@bp.route('/assets')
@login_required
def assets_list():
    status = request.args.get('status', '')
    category = request.args.get('category', '')
    q = Equipment.query
    if status:
        q = q.filter_by(status=status)
    if category:
        q = q.filter_by(category=category)
    items = q.order_by(Equipment.name).all()
    categories = db.session.query(Equipment.category).distinct().filter(Equipment.category.isnot(None), Equipment.category != '').all()
    categories = [c[0] for c in categories]
    return render_template('assets.html', items=items, categories=categories, status=status, category=category)


@bp.route('/assets/new', methods=['GET', 'POST'])
@login_required
@role_required('admin', 'director', 'technician')
def assets_new():
    if request.method == 'POST':
        eq = Equipment(
            name=request.form['name'],
            equipment_type=request.form.get('equipment_type', ''),
            category=request.form.get('category', ''),
            manufacturer=request.form.get('manufacturer', ''),
            model_name=request.form.get('model_name', ''),
            serial_number=request.form.get('serial_number', ''),
            inventory_number=request.form.get('inventory_number', ''),
            year_of_manufacture=safe_int(request.form.get('year_of_manufacture')) or None,
            country_of_origin=request.form.get('country_of_origin', ''),
            voltage=request.form.get('voltage', ''),
            power=request.form.get('power', ''),
            current_rating=request.form.get('current_rating', ''),
            frequency=request.form.get('frequency', ''),
            weight=request.form.get('weight', ''),
            dimensions=request.form.get('dimensions', ''),
            capacity=request.form.get('capacity', ''),
            pressure=request.form.get('pressure', ''),
            temperature_range=request.form.get('temperature_range', ''),
            ip_rating=request.form.get('ip_rating', ''),
            material=request.form.get('material', ''),
            color=request.form.get('color', ''),
            purchase_date=(d := safe_date(request.form.get('purchase_date'))) and d.date() or None,
            purchase_price=safe_float(request.form.get('purchase_price')) or None,
            currency=request.form.get('currency', 'EUR'),
            supplier=request.form.get('supplier', ''),
            invoice_number=request.form.get('invoice_number', ''),
            warranty_start=(d := safe_date(request.form.get('warranty_start'))) and d.date() or None,
            warranty_end=(d := safe_date(request.form.get('warranty_end'))) and d.date() or None,
            warranty_notes=request.form.get('warranty_notes', ''),
            section_id=safe_int(request.form.get('section_id')) or None,
            installation_location=request.form.get('installation_location', ''),
            building=request.form.get('building', ''),
            floor_level=request.form.get('floor_level', ''),
            room=request.form.get('room', ''),
            status=request.form.get('status', 'active'),
            condition=request.form.get('condition', 'good'),
            responsible_person_id=safe_int(request.form.get('responsible_person_id')) or None,
            responsible_user_id=safe_int(request.form.get('responsible_user_id')) or None,
            contractor_id=safe_int(request.form.get('contractor_id')) or None,
            last_service_date=(d := safe_date(request.form.get('last_service_date'))) and d.date() or None,
            next_service_date=(d := safe_date(request.form.get('next_service_date'))) and d.date() or None,
            service_interval_days=safe_int(request.form.get('service_interval_days')) or None,
            maintenance_notes=request.form.get('maintenance_notes', ''),
            description=request.form.get('description', ''),
            notes=request.form.get('notes', ''),
            tags=request.form.get('tags', ''),
        )
        # Photo
        photo = request.files.get('photo')
        if photo and photo.filename:
            eq.photo = save_uploaded_file(photo, 'equipment')
        # Manual
        manual = request.files.get('manual_file')
        if manual and manual.filename:
            eq.manual_file = save_uploaded_file(manual, 'equipment')
        # Certificate
        cert = request.files.get('certificate_file')
        if cert and cert.filename:
            eq.certificate_file = save_uploaded_file(cert, 'equipment')
        db.session.add(eq)
        if not safe_commit():
            flash(_('Save failed'), 'error')
            return redirect(url_for('assets.assets_detail', eq_id=eq.id))
        log_audit('create', 'equipment_asset', eq.id, eq.name)
        flash(_('Equipment added'), 'success')
        return redirect(url_for('assets.assets_detail', eq_id=eq.id))
    sections = FactorySection.query.order_by(FactorySection.name).all()
    verantwoordelijken = Verantwoordelijke.query.order_by(Verantwoordelijke.naam).all()
    users = User.query.filter_by(is_active_user=True).order_by(User.username).all()
    contractors = Contractor.query.order_by(Contractor.company_name).all()
    return render_template('asset_form.html', eq=None, sections=sections, verantwoordelijken=verantwoordelijken, users=users, contractors=contractors)


@bp.route('/assets/<int:eq_id>')
@login_required
def assets_detail(eq_id):
    eq = Equipment.query.get_or_404(eq_id)
    return render_template('asset_detail.html', eq=eq, now=datetime.utcnow())


@bp.route('/assets/<int:eq_id>/edit', methods=['GET', 'POST'])
@login_required
@role_required('admin', 'director', 'technician')
def assets_edit(eq_id):
    eq = Equipment.query.get_or_404(eq_id)
    if request.method == 'POST':
        eq.name = request.form['name']
        eq.equipment_type = request.form.get('equipment_type', '')
        eq.category = request.form.get('category', '')
        eq.manufacturer = request.form.get('manufacturer', '')
        eq.model_name = request.form.get('model_name', '')
        eq.serial_number = request.form.get('serial_number', '')
        eq.inventory_number = request.form.get('inventory_number', '')
        eq.year_of_manufacture = safe_int(request.form.get('year_of_manufacture')) or None
        eq.country_of_origin = request.form.get('country_of_origin', '')
        eq.voltage = request.form.get('voltage', '')
        eq.power = request.form.get('power', '')
        eq.current_rating = request.form.get('current_rating', '')
        eq.frequency = request.form.get('frequency', '')
        eq.weight = request.form.get('weight', '')
        eq.dimensions = request.form.get('dimensions', '')
        eq.capacity = request.form.get('capacity', '')
        eq.pressure = request.form.get('pressure', '')
        eq.temperature_range = request.form.get('temperature_range', '')
        eq.ip_rating = request.form.get('ip_rating', '')
        eq.material = request.form.get('material', '')
        eq.color = request.form.get('color', '')
        eq.purchase_date = (d := safe_date(request.form.get('purchase_date'))) and d.date() or None
        eq.purchase_price = safe_float(request.form.get('purchase_price')) or None
        eq.currency = request.form.get('currency', 'EUR')
        eq.supplier = request.form.get('supplier', '')
        eq.invoice_number = request.form.get('invoice_number', '')
        eq.warranty_start = (d := safe_date(request.form.get('warranty_start'))) and d.date() or None
        eq.warranty_end = (d := safe_date(request.form.get('warranty_end'))) and d.date() or None
        eq.warranty_notes = request.form.get('warranty_notes', '')
        eq.section_id = safe_int(request.form.get('section_id')) or None
        eq.installation_location = request.form.get('installation_location', '')
        eq.building = request.form.get('building', '')
        eq.floor_level = request.form.get('floor_level', '')
        eq.room = request.form.get('room', '')
        eq.status = request.form.get('status', 'active')
        eq.condition = request.form.get('condition', 'good')
        eq.responsible_person_id = safe_int(request.form.get('responsible_person_id')) or None
        eq.responsible_user_id = safe_int(request.form.get('responsible_user_id')) or None
        eq.contractor_id = safe_int(request.form.get('contractor_id')) or None
        eq.last_service_date = (d := safe_date(request.form.get('last_service_date'))) and d.date() or None
        eq.next_service_date = (d := safe_date(request.form.get('next_service_date'))) and d.date() or None
        eq.service_interval_days = safe_int(request.form.get('service_interval_days')) or None
        eq.maintenance_notes = request.form.get('maintenance_notes', '')
        eq.description = request.form.get('description', '')
        eq.notes = request.form.get('notes', '')
        eq.tags = request.form.get('tags', '')
        photo = request.files.get('photo')
        if photo and photo.filename:
            eq.photo = save_uploaded_file(photo, 'equipment')
        manual = request.files.get('manual_file')
        if manual and manual.filename:
            eq.manual_file = save_uploaded_file(manual, 'equipment')
        cert = request.files.get('certificate_file')
        if cert and cert.filename:
            eq.certificate_file = save_uploaded_file(cert, 'equipment')
        if not safe_commit():
            flash(_('Save failed'), 'error')
            return redirect(url_for('assets.assets_detail', eq_id=eq.id))
        flash(_('Equipment updated'), 'success')
        return redirect(url_for('assets.assets_detail', eq_id=eq.id))
    sections = FactorySection.query.order_by(FactorySection.name).all()
    verantwoordelijken = Verantwoordelijke.query.order_by(Verantwoordelijke.naam).all()
    users = User.query.filter_by(is_active_user=True).order_by(User.username).all()
    contractors = Contractor.query.order_by(Contractor.company_name).all()
    return render_template('asset_form.html', eq=eq, sections=sections, verantwoordelijken=verantwoordelijken, users=users, contractors=contractors)


@bp.route('/assets/<int:eq_id>/delete', methods=['POST'])
@login_required
@role_required('admin')
def assets_delete(eq_id):
    eq = Equipment.query.get_or_404(eq_id)
    db.session.delete(eq)
    if not safe_commit():
        flash(_('Save failed'), 'error')
        return redirect(url_for('assets.assets_list'))
    flash(_('Equipment deleted'), 'success')
    return redirect(url_for('assets.assets_list'))


@bp.route('/assets/<int:eq_id>/service', methods=['POST'])
@login_required
@role_required('admin', 'technician')
def assets_add_service(eq_id):
    eq = Equipment.query.get_or_404(eq_id)
    log = EquipmentServiceLog(
        equipment_id=eq.id,
        service_type=request.form.get('service_type', 'maintenance'),
        description=request.form.get('description', ''),
        performed_by=current_user.id,
        cost=safe_float(request.form.get('cost'), 0),
        date=safe_date(request.form.get('date')) or datetime.utcnow(),
        next_date=(d := safe_date(request.form.get('next_date'))) and d.date() or None,
        notes=request.form.get('notes', '')
    )
    if log.next_date:
        eq.next_service_date = log.next_date
    eq.last_service_date = log.date.date() if isinstance(log.date, datetime) else log.date
    db.session.add(log)
    if not safe_commit():
        flash(_('Save failed'), 'error')
        return redirect(url_for('assets.assets_detail', eq_id=eq.id))
    flash(_('Service record added'), 'success')
    return redirect(url_for('assets.assets_detail', eq_id=eq.id))


EQUIPMENT_EXPORT_FIELDS = {
    'name':              lambda e: e.name,
    'equipment_type':    lambda e: e.equipment_type or '',
    'serial_number':     lambda e: e.serial_number or '',
    'year_of_manufacture': lambda e: e.year_of_manufacture or '',
    'manufacturer':      lambda e: e.manufacturer or '',
    'status':            lambda e: e.status or 'active',
    'category':          lambda e: e.category or '',
    'model_name':        lambda e: e.model_name or '',
    'inventory_number':  lambda e: e.inventory_number or '',
    'condition':         lambda e: e.condition or '',
    'installation_location': lambda e: e.installation_location or '',
    'section':           lambda e: e.section.name if e.section else '',
}

EQUIPMENT_EXPORT_FIELD_LABELS = {
    'name': 'Name', 'equipment_type': 'Type', 'serial_number': 'Serial Number',
    'year_of_manufacture': 'Year of Manufacture', 'manufacturer': 'Manufacturer',
    'status': 'Status', 'category': 'Category', 'model_name': 'Model',
    'inventory_number': 'Inventory Number', 'condition': 'Condition',
    'installation_location': 'Location', 'section': 'Section',
}

EQUIPMENT_EXPORT_ALL_FIELDS = ['name', 'equipment_type', 'serial_number', 'manufacturer',
                                'year_of_manufacture', 'status', 'category', 'model_name',
                                'inventory_number', 'condition', 'installation_location', 'section']


@bp.route('/assets/export/<format_type>')
@login_required
@role_required('admin', 'director')
def assets_export(format_type):
    ids_raw = request.args.get('ids', '')
    id_list = [safe_int(x) for x in ids_raw.split(',') if x.strip()]
    if id_list:
        items = Equipment.query.filter(Equipment.id.in_(id_list)).order_by(Equipment.id).all()
    else:
        items = Equipment.query.order_by(Equipment.id).all()

    fields_raw = request.args.get('fields', '')
    field_keys = [f.strip() for f in fields_raw.split(',') if f.strip() in EQUIPMENT_EXPORT_FIELDS] if fields_raw else EQUIPMENT_EXPORT_ALL_FIELDS
    if not field_keys:
        field_keys = EQUIPMENT_EXPORT_ALL_FIELDS

    status_labels = {
        'active': _('Active'), 'maintenance': _('Maintenance'), 'broken': _('Broken'),
        'offline': _('Offline'), 'retired': _('Retired'), 'disposed': _('Disposed'),
    }

    def get_val(e, key):
        if key == 'status':
            return status_labels.get(e.status or 'active', e.status or 'active')
        return EQUIPMENT_EXPORT_FIELDS[key](e)

    header_map = {k: _(EQUIPMENT_EXPORT_FIELD_LABELS[k]) for k in EQUIPMENT_EXPORT_FIELDS}
    headers = [header_map[k] for k in field_keys]

    rows = []
    for e in items:
        rows.append([str(get_val(e, k)) for k in field_keys])

    if format_type == 'xlsx':
        try:
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = 'Equipment'
            ws.append(headers)
            for row in rows:
                ws.append(row)
            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)
            return send_file(buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                             download_name=f'equipment_{datetime.now().strftime("%Y%m%d")}.xlsx', as_attachment=True)
        except ImportError:
            flash(_('Excel export is temporarily unavailable'), 'error')
            return redirect(url_for('assets.assets_list'))
    elif format_type == 'csv':
        import csv
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(headers)
        writer.writerows(rows)
        buf.seek(0)
        return send_file(io.BytesIO(buf.getvalue().encode('utf-8-sig')), mimetype='text/csv',
                         download_name=f'equipment_{datetime.now().strftime("%Y%m%d")}.csv', as_attachment=True)
    elif format_type == 'pdf':
        try:
            fpdf_mod = ensure_fpdf()
            FPDF = fpdf_mod.FPDF
        except Exception:
            flash(_('PDF export is temporarily unavailable'), 'error')
            return redirect(url_for('assets.assets_list'))
        date_str = datetime.now().strftime('%d-%m-%Y %H:%M')
        title = _('Equipment')
        subtitle = f'{date_str} - {len(items)} {_("items")}'
        footer_txt = _('Generated by ProMaster CRM')
        page_txt = _('page')

        font_path = find_pdf_font()
        unicode_font = bool(font_path)

        def pdf_text(s):
            s = str(s if s is not None else '')
            if unicode_font:
                return s
            return s.encode('latin-1', 'replace').decode('latin-1')

        try:
            class EquipPDF(FPDF):
                def header(self):
                    fn = 'AppFont' if unicode_font else 'Helvetica'
                    self.set_font(fn, 'B', 12)
                    self.cell(0, 8, pdf_text(f'{title} - ProMaster'), new_x='LMARGIN', new_y='NEXT')
                    self.set_font(fn, '', 8)
                    self.cell(0, 5, pdf_text(subtitle), new_x='LMARGIN', new_y='NEXT')
                    self.ln(2)

                def footer(self):
                    fn = 'AppFont' if unicode_font else 'Helvetica'
                    self.set_y(-12)
                    self.set_font(fn, '', 7)
                    self.set_text_color(128)
                    self.cell(0, 8, pdf_text(f'{footer_txt} - {page_txt} {self.page_no()}/{{nb}}'), align='C')

            pdf = EquipPDF(orientation='L', unit='mm', format='A4')
            pdf.alias_nb_pages()
            pdf.set_auto_page_break(auto=True, margin=15)
            if unicode_font:
                try:
                    pdf.add_font('AppFont', '', font_path)
                    pdf.add_font('AppFont', 'B', font_path)
                    base_font = 'AppFont'
                except Exception:
                    unicode_font = False
                    base_font = 'Helvetica'
            else:
                base_font = 'Helvetica'
            pdf.add_page()

            n_cols = len(headers)
            col_w = [277 / n_cols] * n_cols
            pdf.set_font(base_font, 'B', 7)
            pdf.set_fill_color(44, 62, 80)
            pdf.set_text_color(255)
            for i, h in enumerate(headers):
                pdf.cell(col_w[i], 7, pdf_text(h), border=1, fill=True)
            pdf.ln()
            pdf.set_font(base_font, '', 6.5)
            pdf.set_text_color(51)
            for idx, row in enumerate(rows):
                fill = (idx % 2 == 1)
                if fill:
                    pdf.set_fill_color(249, 249, 249)
                for i, val in enumerate(row):
                    pdf.cell(col_w[i], 6, pdf_text(val), border=1, fill=fill)
                pdf.ln()

            buf = io.BytesIO()
            pdf.output(buf)
            buf.seek(0)
            return send_file(buf, mimetype='application/pdf',
                             download_name=f'equipment_{datetime.now().strftime("%Y%m%d")}.pdf', as_attachment=True)
        except Exception as e:
            flash(_('PDF export failed: %(err)s', err=str(e)), 'error')
            return redirect(url_for('assets.assets_list'))
    flash(_('Unsupported format'), 'error')
    return redirect(url_for('assets.assets_list'))
