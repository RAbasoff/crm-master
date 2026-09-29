"""
repairs blueprint
"""
from datetime import datetime, timedelta
from flask import (Blueprint, render_template, request, redirect, url_for, flash,
                   jsonify, send_file, session, g)
from flask_login import login_required, current_user
from flask_babel import gettext as _
from werkzeug.utils import secure_filename
import os, io, json

from models import (db, EquipmentRepair, GasSystemComponent)
from utils import role_required, safe_commit, safe_date, safe_int, ensure_fpdf, find_pdf_font

bp = Blueprint('repairs', __name__)

@bp.route('/repairs')
@login_required
@role_required('admin', 'director', 'technician')
def repairs_list():
    gas_type = request.args.get('gas', '')
    status = request.args.get('status', '')
    q = EquipmentRepair.query
    if status:
        q = q.filter_by(status=status)
    if gas_type:
        q = q.join(GasSystemComponent).filter(GasSystemComponent.gas_type == gas_type)
    repairs = q.order_by(EquipmentRepair.date_broken.desc()).all()
    components = GasSystemComponent.query.order_by(GasSystemComponent.gas_type, GasSystemComponent.component_type).all()
    return render_template('repairs.html', repairs=repairs, components=components, gas_filter=gas_type, status_filter=status)


@bp.route('/repairs/new', methods=['GET', 'POST'])
@login_required
@role_required('admin', 'technician')
def repair_new():
    if request.method == 'POST':
        comp_id = request.form.get('component_id')
        if not comp_id:
            flash(_('Select component'), 'error')
            return redirect(url_for('repairs.repair_new'))
        comp_id = int(comp_id)
        comp = GasSystemComponent.query.get(comp_id)
        cost_str = request.form.get('repair_cost', '').strip()
        try:
            date_broken = datetime.strptime(request.form['date_broken'], '%Y-%m-%dT%H:%M')
        except (ValueError, KeyError):
            flash(_('Invalid date format'), 'error')
            return redirect(url_for('repairs.repair_new'))
        r = EquipmentRepair(
            component_id=comp_id,
            fault_description=request.form['fault_description'],
            date_broken=date_broken,
            repair_company=request.form.get('repair_company', ''),
            repair_description=request.form.get('repair_description', ''),
            repair_cost=float(cost_str) if cost_str else 0,
            date_sent=safe_date(request.form.get('date_sent'), '%Y-%m-%dT%H:%M'),
            date_repaired=safe_date(request.form.get('date_repaired'), '%Y-%m-%dT%H:%M'),
            date_installed=safe_date(request.form.get('date_installed'), '%Y-%m-%dT%H:%M'),
            status=request.form.get('status', 'broken'),
            notes=request.form.get('notes', ''),
            created_by=current_user.id
        )
        # Update component status
        if comp:
            if r.status == 'broken':
                comp.status = 'faulty'
            elif r.status in ('in_repair', 'repaired'):
                comp.status = 'replaced'
            elif r.status == 'installed':
                comp.status = 'ok'
                comp.installed_at = r.date_installed
        db.session.add(r)
        if not safe_commit():
            flash(_('Save failed'), 'error')
            return redirect(url_for('repairs.repairs_list'))
        flash(_('Repair record created'), 'success')
        return redirect(url_for('repairs.repairs_list'))
    components = GasSystemComponent.query.order_by(GasSystemComponent.gas_type, GasSystemComponent.component_type).all()
    return render_template('repair_form.html', repair=None, components=components)


@bp.route('/repairs/<int:repair_id>/edit', methods=['GET', 'POST'])
@login_required
@role_required('admin', 'technician')
def repair_edit(repair_id):
    r = EquipmentRepair.query.get_or_404(repair_id)
    if request.method == 'POST':
        try:
            r.date_broken = datetime.strptime(request.form['date_broken'], '%Y-%m-%dT%H:%M')
        except (ValueError, KeyError):
            flash(_('Invalid date format'), 'error')
            return redirect(url_for('repairs.repair_edit', repair_id=r.id))
        r.fault_description = request.form['fault_description']
        r.repair_company = request.form.get('repair_company', '')
        r.repair_description = request.form.get('repair_description', '')
        cost_str = request.form.get('repair_cost', '').strip()
        r.repair_cost = float(cost_str) if cost_str else 0
        r.date_sent = safe_date(request.form.get('date_sent'), '%Y-%m-%dT%H:%M')
        r.date_repaired = safe_date(request.form.get('date_repaired'), '%Y-%m-%dT%H:%M')
        r.date_installed = safe_date(request.form.get('date_installed'), '%Y-%m-%dT%H:%M')
        r.status = request.form.get('status', r.status)
        r.notes = request.form.get('notes', '')
        # Update component
        comp = r.component
        if comp:
            if r.status == 'installed':
                comp.status = 'ok'
                comp.installed_at = r.date_installed
            elif r.status == 'broken':
                comp.status = 'faulty'
        if not safe_commit():
            flash(_('Save failed'), 'error')
            return redirect(url_for('repairs.repairs_list'))
        flash(_('Repair record updated'), 'success')
        return redirect(url_for('repairs.repairs_list'))
    components = GasSystemComponent.query.order_by(GasSystemComponent.gas_type, GasSystemComponent.component_type).all()
    return render_template('repair_form.html', repair=r, components=components)


@bp.route('/repairs/<int:repair_id>/status', methods=['POST'])
@login_required
@role_required('admin', 'technician')
def repair_status(repair_id):
    r = EquipmentRepair.query.get_or_404(repair_id)
    new_status = request.form.get('status')
    if new_status in ('broken', 'in_repair', 'repaired', 'installed'):
        r.status = new_status
        comp = r.component
        if new_status == 'in_repair':
            r.date_sent = datetime.utcnow()
            if comp: comp.status = 'replaced'
        elif new_status == 'repaired':
            r.date_repaired = datetime.utcnow()
            if comp: comp.status = 'replaced'
        elif new_status == 'installed':
            r.date_installed = datetime.utcnow()
            if comp:
                comp.status = 'ok'
                comp.installed_at = r.date_installed
        if not safe_commit():
            flash(_('Save failed'), 'error')
            return redirect(url_for('repairs.repairs_list'))
    flash(_('Status updated'), 'success')
    return redirect(url_for('repairs.repairs_list'))


@bp.route('/repairs/stats')
@login_required
@role_required('admin', 'director', 'technician')
def repairs_stats():
    all_repairs = EquipmentRepair.query.order_by(EquipmentRepair.date_broken.desc()).all()
    # Stats
    total = len(all_repairs)
    by_component = {}
    by_company = {}
    total_cost = 0
    total_days = 0
    days_count = 0
    for r in all_repairs:
        # By component type
        ctype = r.component.component_type if r.component else 'unknown'
        if ctype not in by_component:
            by_component[ctype] = {'count': 0, 'cost': 0, 'days': []}
        by_component[ctype]['count'] += 1
        by_component[ctype]['cost'] += r.repair_cost or 0
        # Days in repair
        if r.date_broken and r.date_installed:
            days = (r.date_installed - r.date_broken).days
            by_component[ctype]['days'].append(days)
            total_days += days
            days_count += 1
        # By company
        if r.repair_company:
            if r.repair_company not in by_company:
                by_company[r.repair_company] = {'count': 0, 'cost': 0}
            by_company[r.repair_company]['count'] += 1
            by_company[r.repair_company]['cost'] += r.repair_cost or 0
        total_cost += r.repair_cost or 0
    # Average days
    avg_days = round(total_days / days_count, 1) if days_count else 0
    # Add averages to by_component
    for ctype in by_component:
        days_list = by_component[ctype]['days']
        by_component[ctype]['avg_days'] = round(sum(days_list) / len(days_list), 1) if days_list else 0
    return render_template('repairs_stats.html',
        total=total, total_cost=total_cost, avg_days=avg_days,
        by_component=by_component, by_company=by_company,
        recent=all_repairs[:20])


REPAIR_EXPORT_FIELDS = {
    'id':                lambda r: r.id,
    'component':         lambda r: r.component.name if r.component else '',
    'component_type':    lambda r: r.component.component_type if r.component else '',
    'gas_type':          lambda r: r.component.gas_type if r.component else '',
    'fault_description': lambda r: r.fault_description or '',
    'date_broken':       lambda r: r.date_broken.strftime('%d-%m-%Y %H:%M') if r.date_broken else '',
    'repair_company':    lambda r: r.repair_company or '',
    'repair_description': lambda r: r.repair_description or '',
    'repair_cost':       lambda r: r.repair_cost or 0,
    'date_sent':         lambda r: r.date_sent.strftime('%d-%m-%Y %H:%M') if r.date_sent else '',
    'date_repaired':     lambda r: r.date_repaired.strftime('%d-%m-%Y %H:%M') if r.date_repaired else '',
    'date_installed':    lambda r: r.date_installed.strftime('%d-%m-%Y %H:%M') if r.date_installed else '',
    'status':            lambda r: r.status or '',
    'notes':             lambda r: r.notes or '',
}

REPAIR_EXPORT_FIELD_LABELS = {
    'id': 'ID', 'component': 'Component', 'component_type': 'Component Type',
    'gas_type': 'Gas', 'fault_description': 'Fault', 'date_broken': 'Date Broken',
    'repair_company': 'Company', 'repair_description': 'Repair Description',
    'repair_cost': 'Cost', 'date_sent': 'Date Sent', 'date_repaired': 'Date Repaired',
    'date_installed': 'Date Installed', 'status': 'Status', 'notes': 'Notes',
}

REPAIR_EXPORT_ALL_FIELDS = ['id', 'component', 'gas_type', 'fault_description', 'date_broken',
                             'repair_company', 'repair_cost', 'date_sent', 'date_repaired',
                             'date_installed', 'status']


@bp.route('/repairs/export/<format_type>')
@login_required
@role_required('admin', 'director', 'technician')
def repairs_export(format_type):
    ids_raw = request.args.get('ids', '')
    id_list = [safe_int(x) for x in ids_raw.split(',') if x.strip()]
    if id_list:
        repairs = EquipmentRepair.query.filter(EquipmentRepair.id.in_(id_list)).order_by(EquipmentRepair.id).all()
    else:
        repairs = EquipmentRepair.query.order_by(EquipmentRepair.date_broken.desc()).all()

    fields_raw = request.args.get('fields', '')
    field_keys = [f.strip() for f in fields_raw.split(',') if f.strip() in REPAIR_EXPORT_FIELDS] if fields_raw else REPAIR_EXPORT_ALL_FIELDS
    if not field_keys:
        field_keys = REPAIR_EXPORT_ALL_FIELDS

    status_labels = {
        'broken': _('Broken'), 'in_repair': _('In repair'),
        'repaired': _('Repaired'), 'installed': _('Installed'),
    }

    def get_val(r, key):
        if key == 'status':
            return status_labels.get(r.status or 'broken', r.status or '')
        return REPAIR_EXPORT_FIELDS[key](r)

    header_map = {k: _(REPAIR_EXPORT_FIELD_LABELS[k]) for k in REPAIR_EXPORT_FIELDS}
    headers = [header_map[k] for k in field_keys]
    rows = []
    for r in repairs:
        rows.append([str(get_val(r, k)) for k in field_keys])

    if format_type == 'xlsx':
        try:
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = 'Repairs'
            ws.append(headers)
            for row in rows:
                ws.append(row)
            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)
            return send_file(buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                             download_name=f'repairs_{datetime.now().strftime("%Y%m%d")}.xlsx', as_attachment=True)
        except ImportError:
            flash(_('Excel export is temporarily unavailable'), 'error')
            return redirect(url_for('repairs.repairs_list'))
    elif format_type == 'csv':
        import csv
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(headers)
        writer.writerows(rows)
        buf.seek(0)
        return send_file(io.BytesIO(buf.getvalue().encode('utf-8-sig')), mimetype='text/csv',
                         download_name=f'repairs_{datetime.now().strftime("%Y%m%d")}.csv', as_attachment=True)
    elif format_type == 'pdf':
        try:
            fpdf_mod = ensure_fpdf()
            FPDF = fpdf_mod.FPDF
        except Exception:
            flash(_('PDF export is temporarily unavailable'), 'error')
            return redirect(url_for('repairs.repairs_list'))
        date_str = datetime.now().strftime('%d-%m-%Y %H:%M')
        title = _('Equipment Repairs')
        subtitle = f'{date_str} - {len(repairs)} {_("records")}'
        footer_txt = _('Generated by ProMaster CRM')
        page_txt = _('page')
        font_path = find_pdf_font()
        unicode_font = bool(font_path)

        def pdf_text(s):
            s = str(s if s is not None else '')
            return s if unicode_font else s.encode('latin-1', 'replace').decode('latin-1')

        try:
            class RepairPDF(FPDF):
                def header(self):
                    fn = 'AppFont' if unicode_font else 'Helvetica'
                    self.set_font(fn, 'B', 12)
                    self.cell(0, 8, pdf_text(f'{title} - ProMaster'), new_x='LMARGIN', new_y='NEXT')
                    self.set_font(fn, '', 8)
                    self.cell(0, 5, pdf_text(subtitle), new_x='LMARGIN', new_y='NEXT')
                    self.ln(2)
                def footer(self):
                    fn = 'AppFont' if unicode_font else 'Helvetica'
                    self.set_y(-12)
                    self.set_font(fn, '', 7)
                    self.set_text_color(128)
                    self.cell(0, 8, pdf_text(f'{footer_txt} - {page_txt} {self.page_no()}/{{nb}}'), align='C')

            pdf = RepairPDF(orientation='L', unit='mm', format='A4')
            pdf.alias_nb_pages()
            pdf.set_auto_page_break(auto=True, margin=15)
            if unicode_font:
                try:
                    pdf.add_font('AppFont', '', font_path)
                    pdf.add_font('AppFont', 'B', font_path)
                    base_font = 'AppFont'
                except Exception:
                    unicode_font = False
                    base_font = 'Helvetica'
            else:
                base_font = 'Helvetica'
            pdf.add_page()
            n_cols = len(headers)
            col_w = [277 / n_cols] * n_cols
            pdf.set_font(base_font, 'B', 7)
            pdf.set_fill_color(44, 62, 80)
            pdf.set_text_color(255)
            for i, h in enumerate(headers):
                pdf.cell(col_w[i], 7, pdf_text(h), border=1, fill=True)
            pdf.ln()
            pdf.set_font(base_font, '', 6.5)
            pdf.set_text_color(51)
            for idx, row in enumerate(rows):
                fill = (idx % 2 == 1)
                if fill:
                    pdf.set_fill_color(249, 249, 249)
                for i, val in enumerate(row):
                    pdf.cell(col_w[i], 6, pdf_text(val), border=1, fill=fill)
                pdf.ln()
            buf = io.BytesIO()
            pdf.output(buf)
            buf.seek(0)
            return send_file(buf, mimetype='application/pdf',
                             download_name=f'repairs_{datetime.now().strftime("%Y%m%d")}.pdf', as_attachment=True)
        except Exception as e:
            flash(_('PDF export failed: %(err)s', err=str(e)), 'error')
            return redirect(url_for('repairs.repairs_list'))
    flash(_('Unsupported format'), 'error')
    return redirect(url_for('repairs.repairs_list'))
